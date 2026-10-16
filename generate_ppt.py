from pptx import Presentation
import os

prs = Presentation()

# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[0])
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "RetailIQ: End-to-End E-Commerce Analytics & Customer Churn MLOps"
subtitle.text = "From Raw Data Ingestion to Real-Time Predictive AI"

def add_bullet_slide(prs, title_text, bullet_points):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = title_text
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    for i, pt in enumerate(bullet_points):
        if i == 0:
            p = tf.paragraphs[0]
            p.text = pt
        else:
            p = tf.add_paragraph()
            p.text = pt
            p.level = 0
            
# Slide 2
add_bullet_slide(prs, "Project Vision & Objectives", [
    "Problem: E-commerce platforms struggle to identify at-risk customers.",
    "Objective: Build an automated, scalable pipeline that transforms raw data into actionable BI and real-time AI.",
    "Phase 1: Robust ETL, Data Warehouse, and Airflow orchestration.",
    "Phase 2: Feature Engineering, MLflow tracking, FastAPI deployment, Streamlit UI."
])

# Slide 3
add_bullet_slide(prs, "Architecture Overview", [
    "Orchestration: Apache Airflow",
    "Storage / Warehouse: PostgreSQL (Star Schema)",
    "ML Experimentation: MLflow & Scikit-Learn",
    "Model Deployment: FastAPI",
    "Front-End Dashboard: Streamlit",
    "Infrastructure: Docker & Docker Compose"
])

# Slide 4
add_bullet_slide(prs, "Phase 1 - Ingestion & Quality", [
    "Data Source: +500k e-commerce transactions across 13 monthly batches.",
    "Quarantine Framework: Enforced strict business rules (valid prices, etc).",
    "Results: Achieved a 99.53% Validation Pass Rate.",
    "Auditability: 2,500+ records safely quarantined with explicit error codes."
])

# Slide 5
add_bullet_slide(prs, "Phase 1 - Data Warehouse", [
    "Dimensional Modeling: Star Schema with core dimensions and fact_sales.",
    "Sales Performance Mart: Revenue, AOV, and volume metrics.",
    "RFM Segmentation Mart: Recency, Frequency, Monetary scores for Loyalty Tiers."
])

# Slide 6
add_bullet_slide(prs, "Phase 2 - Feature Engineering", [
    "Objective: Predict if an active customer will churn.",
    "Temporal Splitting (Zero Leakage):",
    "Observation Window (180 days): 14 behavioral features.",
    "Target Window (90 days): Binary churn labels based on inactivity.",
    "Result: Predictive feature set for 2,786 unique customers."
])

# Slide 7
add_bullet_slide(prs, "Phase 2 - Model Training & MLOps", [
    "Algorithm Benchmarking: Logistic Regression, Random Forest, Gradient Boosting.",
    "MLflow Registry: Tracked hyper-parameters, metrics, and artifacts.",
    "The Winner: Gradient Boosting Classifier (ROC-AUC: 0.6950).",
    "Deployment Ready: Serialized as production_churn_model.pkl."
])

# Slide 8
add_bullet_slide(prs, "Phase 2 - Production Deployment", [
    "Microservices Architecture: Deployed the winning model in a FastAPI REST service.",
    "Docker Compose Network:",
    " - ecommerce_postgres (Data Warehouse)",
    " - ecommerce_fastapi_api (AI Backend)",
    " - ecommerce_streamlit_app (Front-End)"
])

# Slide 9
add_bullet_slide(prs, "The Final Product", [
    "Interactive Streamlit Dashboard: 8 Analytical Tabs.",
    "Executive Revenue Trends & Global Heatmaps.",
    "3D RFM Behavioral Scatter Plots.",
    "Real-Time Churn Portal: Instantly receive Churn Probability and Recommendations."
])

# Slide 10
add_bullet_slide(prs, "Conclusion & Next Steps", [
    "Business Impact: Shifted from retrospective reporting to proactive AI strategies.",
    "Next Steps:",
    " - Implement automated model retraining DAGs in Airflow.",
    " - Integrate A/B testing for retention playbooks."
])

prs.save("docs/RetailIQ_Presentation.pptx")
print("Presentation generated successfully!")
